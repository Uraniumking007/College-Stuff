#!/usr/bin/env python3
"""Rebuild Practical 6-11 docx from Template.docx with Aim, monospace code, Remix screenshots."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Emu, Twips
from docx.enum.text import WD_LINE_SPACING

ROOT = Path(".")
PRACT = ROOT / "Practicals"
SHOT = PRACT / "screenshots" / "remix"
TPL = PRACT / "Template.docx"

AIMS = {
    6: "To use Remix IDE to write, compile, and deploy a SimpleStorage contract on Remix VM and test the set and get functions.",
    7: "To demonstrate Solidity data types (bool, int, uint, address, string, bytes32) by compiling and deploying DataTypesDemo and reading the stored values in Remix.",
    8: "To implement an owner-restricted Counter contract, deploy it on Remix VM, and read the count.",
    9: "To implement a simple voting contract, deploy it on Remix VM, add a candidate, and cast a vote.",
    10: "To implement a SimpleWallet contract that accepts ETH and deploy it on Remix VM.",
    11: "To implement a StudentRecordSystem with a registration fee and deploy it on Remix VM.",
}

SHOT_KEYS = ["code", "compile", "deploy", "interact"]

def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)

def add_para(doc, text, *, bold=False, mono=False, size=12, space_after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Courier New" if mono else "Times New Roman"
    rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    face = "Courier New" if mono else "Times New Roman"
    rFonts.set(qn("w:ascii"), face)
    rFonts.set(qn("w:hAnsi"), face)
    return p

def add_picture(doc, path):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.3))

def build(n):
    doc = Document(str(TPL))
    header = doc.sections[0].header.paragraphs[0].text
    assert "240093116002" in header, header
    clear_body(doc)
    add_para(doc, f"Practical No.{n}", bold=True, size=14, space_after=8)
    add_para(doc, "Aim: " + AIMS[n], size=12, space_after=8)
    add_para(doc, "Code:", bold=True, size=12, space_after=4)
    src = (ROOT / "contracts").glob("*.sol")
    names = {
        6: "SimpleStorage.sol", 7: "DataTypesDemo.sol", 8: "Counter.sol",
        9: "Voting.sol", 10: "SimpleWallet.sol", 11: "StudentRecordSystem.sol",
    }
    code = (ROOT / "contracts" / names[n]).read_text()
    for line in code.splitlines():
        add_para(doc, line if line else " ", mono=True, size=8, space_after=0)
    add_para(doc, "", size=8, space_after=4)
    add_para(doc, "Output:", bold=True, size=12, space_after=4)
    used = []
    missing = []
    for key in SHOT_KEYS:
        path = SHOT / f"p{n}_{key}.png"
        if not path.exists():
            missing.append(key)
            continue
        add_picture(doc, path)
        used.append(path.name)
    if missing:
        add_para(doc, "Not captured: " + ", ".join(missing), size=11)
    out = PRACT / f"Practical {n}.docx"
    doc.save(str(out))
    # header still there
    check = Document(str(out))
    h = check.sections[0].header.paragraphs[0].text
    print(n, "header", h[:40], "shots", used, "missing", missing)

if __name__ == "__main__":
    for n in range(6, 12):
        build(n)
